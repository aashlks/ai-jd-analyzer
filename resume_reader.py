"""Extract editable text from a PDF or DOCX resume in memory."""

import io
import re

from docx import Document
from pypdf import PdfReader


MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_PDF_PAGES = 8
MAX_TEXT_CHARS = 16000


class ResumeReadError(ValueError):
    """The uploaded resume cannot be used as text."""


def extract_resume_text(filename: str, file_bytes: bytes) -> str:
    """Read the file without saving it to disk or calling an external API."""

    if not file_bytes:
        raise ResumeReadError("简历文件是空的。")
    if len(file_bytes) > MAX_FILE_BYTES:
        raise ResumeReadError("简历文件不能超过 5 MB。")

    extension = filename.lower().rsplit(".", 1)[-1]
    if extension not in ("pdf", "docx"):
        raise ResumeReadError("目前只支持 PDF 和 DOCX 简历。")

    try:
        if extension == "pdf":
            reader = PdfReader(io.BytesIO(file_bytes), strict=False)
            if reader.is_encrypted:
                raise ResumeReadError("暂不支持加密的 PDF，请先导出未加密版本。")
            if len(reader.pages) > MAX_PDF_PAGES:
                raise ResumeReadError("PDF 最多支持 8 页，请上传简历而不是整份作品集。")
            parts = [page.extract_text() or "" for page in reader.pages]
        else:
            document = Document(io.BytesIO(file_bytes))
            parts = [paragraph.text for paragraph in document.paragraphs]
            # 很多简历把经历放在表格中，因此也读出单元格文字。
            for table in document.tables:
                for row in table.rows:
                    parts.extend(cell.text for cell in row.cells)
    except ResumeReadError:
        raise
    except Exception as error:
        # 文件解析器可能抛出多种损坏/格式错误；不要显示内部堆栈或文件内容。
        raise ResumeReadError("无法读取简历。请确认文件未损坏，且不是加密文件。") from error

    text = "\n".join(line.strip() for part in parts for line in part.splitlines() if line.strip())
    if len(text) < 20:
        raise ResumeReadError("几乎没有提取到文字。扫描版 PDF 暂不支持，请上传可选中文字的 PDF 或 DOCX。")
    if len(text) > MAX_TEXT_CHARS:
        raise ResumeReadError("简历文字过长，请精简到约 16000 字以内。")
    return text


def redact_basic_contacts(text: str) -> str:
    """Hide common contact formats; the user must still review the preview."""

    text = re.sub(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}", "[邮箱已隐藏]", text)
    text = re.sub(r"(?<!\d)1[3-9]\d{9}(?!\d)", "[手机号已隐藏]", text)
    text = re.sub(r"(?<!\d)\d{17}[\dXx](?!\d)", "[证件号码已隐藏]", text)
    return text
