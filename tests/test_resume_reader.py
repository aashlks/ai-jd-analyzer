"""Offline tests for reading a resume without sending it to a service."""

import io
import unittest

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from resume_reader import ResumeReadError, extract_resume_text, redact_basic_contacts


class ResumeReaderTests(unittest.TestCase):
    def test_extracts_docx_paragraphs_and_tables(self) -> None:
        document = Document()
        document.add_paragraph("AI internship applicant with Python projects")
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "Built a small RAG application"
        output = io.BytesIO()
        document.save(output)

        text = extract_resume_text("resume.docx", output.getvalue())

        self.assertIn("Python projects", text)
        self.assertIn("RAG application", text)

    def test_reports_scanned_or_empty_pdf(self) -> None:
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        output = io.BytesIO()
        writer.write(output)

        with self.assertRaisesRegex(ResumeReadError, "没有提取到文字"):
            extract_resume_text("resume.pdf", output.getvalue())

    def test_extracts_selectable_pdf_text(self) -> None:
        writer = PdfWriter()
        page = writer.add_blank_page(width=300, height=200)
        page[NameObject("/Resources")] = DictionaryObject(
            {
                NameObject("/Font"): DictionaryObject(
                    {
                        NameObject("/F1"): DictionaryObject(
                            {
                                NameObject("/Type"): NameObject("/Font"),
                                NameObject("/Subtype"): NameObject("/Type1"),
                                NameObject("/BaseFont"): NameObject("/Helvetica"),
                            }
                        )
                    }
                )
            }
        )
        content = DecodedStreamObject()
        content.set_data(
            b"BT /F1 12 Tf 10 100 Td (Python AI internship resume experience) Tj ET"
        )
        page[NameObject("/Contents")] = writer._add_object(content)
        output = io.BytesIO()
        writer.write(output)

        text = extract_resume_text("resume.pdf", output.getvalue())

        self.assertIn("Python AI internship resume experience", text)

    def test_rejects_other_file_types(self) -> None:
        with self.assertRaisesRegex(ResumeReadError, "只支持 PDF 和 DOCX"):
            extract_resume_text("resume.txt", b"AI internship applicant with Python projects")

    def test_masks_common_contact_formats(self) -> None:
        text = redact_basic_contacts(
            "邮箱 abc@example.com，手机 13812345678，证件 110101200001011234"
        )

        self.assertNotIn("abc@example.com", text)
        self.assertNotIn("13812345678", text)
        self.assertNotIn("110101200001011234", text)


if __name__ == "__main__":
    unittest.main()
